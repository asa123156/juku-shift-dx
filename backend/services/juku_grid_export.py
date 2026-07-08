"""SQLite の割当データ（正本）を拠点テンプレートに変換して Excel 出力する。"""

from __future__ import annotations

import io
import re
from datetime import date, datetime, time
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook

from services.assignment_store import get_assignments_for_date
from services.entity_store import get_student
from services.juku_grid_parser import list_schedule_sheet_names, resolve_dates_for_sheet
from services.location_config import load_grid_map, output_dir_for_period, resolve_location_slug, template_path
from services.period_store import get_period, open_dates_for_period
from services.slot_timing import generate_time_slots


def _slot_to_time(slot: int, location_slug: str | None) -> time:
    for item in generate_time_slots():
        if item["slot"] == slot:
            hh, mm = item["start"].split(":")
            return time(int(hh), int(mm))
    raise HTTPException(status_code=400, detail=f"不明なコマ: {slot}")


def _count_names_in_row(ws, row_idx: int) -> int:
    return sum(1 for col in range(1, ws.max_column + 1) if ws.cell(row=row_idx, column=col).value == "氏名")


def _find_header_row_for_slot(ws, slot: int, location_slug: str | None) -> int | None:
    cfg = load_grid_map(location_slug)
    time_col = cfg["time_column"] + 1
    target_time = _slot_to_time(slot, location_slug)

    for row_idx in range(1, ws.max_row + 1):
        cell_val = ws.cell(row=row_idx, column=time_col).value
        if not isinstance(cell_val, time):
            continue
        if cell_val.hour != target_time.hour or cell_val.minute != target_time.minute:
            continue
        if _count_names_in_row(ws, row_idx) >= 2:
            return row_idx
        if row_idx + 1 <= ws.max_row and _count_names_in_row(ws, row_idx + 1) >= 2:
            return row_idx + 1

    return None


def _data_row_range(ws, header_row: int, slot: int, location_slug: str | None) -> range:
    cfg = load_grid_map(location_slug)
    time_col = cfg["time_column"] + 1
    target_time = _slot_to_time(slot, location_slug)
    end_row = ws.max_row
    for row_idx in range(header_row + 1, ws.max_row + 1):
        val = ws.cell(row=row_idx, column=time_col).value
        if isinstance(val, time) and not (val.hour == target_time.hour and val.minute == target_time.minute):
            end_row = row_idx - 1
            break
        if _count_names_in_row(ws, row_idx) >= 2 and row_idx != header_row:
            end_row = row_idx - 1
            break
    return range(header_row + 1, end_row + 1)


def _teacher_col_index(teacher_columns: list[int], slot: int, teacher_id: int, assignments: list[dict]) -> int:
    teacher_ids = sorted({a["teacher_id"] for a in assignments if a["slot"] == slot})
    if teacher_id not in teacher_ids:
        return teacher_columns[0]
    idx = teacher_ids.index(teacher_id)
    if idx >= len(teacher_columns):
        idx = len(teacher_columns) - 1
    return teacher_columns[idx]


def _write_cell(ws, row: int, col: int, value) -> None:
    from openpyxl.cell.cell import MergedCell

    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for merged_range in list(ws.merged_cells.ranges):
            if (
                merged_range.min_row <= row <= merged_range.max_row
                and merged_range.min_col <= col <= merged_range.max_col
            ):
                ws.unmerge_cells(str(merged_range))
                break
    ws.cell(row=row, column=col, value=value)


def _set_date_header(ws, iso_date: str, location_slug: str | None) -> None:
    cfg = load_grid_map(location_slug)
    dr, dc = cfg["date_header_row"], cfg["date_header_col"]
    _write_cell(
        ws,
        dr + 1,
        dc + 1,
        datetime.combine(date.fromisoformat(iso_date), time.min),
    )


def _fill_assignments_on_sheet(
    ws,
    iso_date: str,
    assignments: list[dict],
    location_slug: str | None,
) -> None:
    cfg = load_grid_map(location_slug)
    teacher_columns: list[int] = cfg["teacher_columns"]
    offsets = cfg["offsets_from_teacher"]

    if not assignments:
        return

    fill_counts: dict[tuple[int, int], int] = {}
    for assign in sorted(assignments, key=lambda x: (x["slot"], x["teacher_id"], x["student_id"])):
        slot = int(assign["slot"])
        header_row = _find_header_row_for_slot(ws, slot, location_slug)
        if header_row is None:
            continue
        teacher_col = _teacher_col_index(teacher_columns, slot, assign["teacher_id"], assignments)
        tc = teacher_col + 1

        count_key = (slot, assign["teacher_id"])
        fill_idx = fill_counts.get(count_key, 0)
        fill_counts[count_key] = fill_idx + 1

        use_b = fill_idx % 2 == 1
        name_key = "name_b" if use_b else "name_a"
        subject_key = "subject_b" if use_b else "subject_a"
        grade_key = "grade_b" if use_b else "grade_a"

        data_rows = list(_data_row_range(ws, header_row, slot, location_slug))
        if not data_rows:
            continue
        target_row = data_rows[min(fill_idx // 2, len(data_rows) - 1)]

        name_col = tc + offsets[name_key]
        subject_col = tc + offsets[subject_key]
        grade_col = tc + offsets[grade_key]
        teacher_cell_col = tc + offsets["teacher"]

        try:
            student = get_student(assign["student_id"])
            g_label = student["grade_label"]
        except HTTPException:
            g_label = ""

        _write_cell(ws, target_row, name_col, assign["student_name"])
        _write_cell(ws, target_row, subject_col, assign["subject"])
        if g_label:
            _write_cell(ws, target_row, name_col - 1, g_label)
            _write_cell(ws, target_row, grade_col, g_label)
        _write_cell(ws, target_row, teacher_cell_col, assign["teacher_name"])


def _workbook_bytes(wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def _daily_template_sheet_name(wb) -> str:
    daily = [s for s in wb.sheetnames if re.match(r"\d+月\d+日", s)]
    if not daily:
        raise HTTPException(status_code=500, detail="日次シートがテンプレートにありません")
    return daily[0]


def _sheet_title_for_date(iso_date: str) -> str:
    d = date.fromisoformat(iso_date)
    return f"{d.month}月{d.day}日"


def _build_workbook_for_open_dates(open_dates: list[str], location_slug: str) -> bytes:
    slug = resolve_location_slug(location_slug)
    if not open_dates:
        raise HTTPException(status_code=404, detail="開校日がありません")

    template_bytes = template_path(slug).read_bytes()
    wb = load_workbook(io.BytesIO(template_bytes))
    template_name = _daily_template_sheet_name(wb)
    template_ws = wb[template_name]
    for name in list(wb.sheetnames):
        if name != template_name:
            del wb[name]

    worksheets = [template_ws]
    for _ in range(1, len(open_dates)):
        worksheets.append(wb.copy_worksheet(template_ws))

    for ws, iso_date in zip(worksheets, open_dates):
        ws.title = _sheet_title_for_date(iso_date)[:31]
        assignments = get_assignments_for_date(iso_date)
        _set_date_header(ws, iso_date, slug)
        _fill_assignments_on_sheet(ws, iso_date, assignments, slug)

    return _workbook_bytes(wb)


def build_schedule_workbook_from_db(period_id: int, location_slug: str | None = None) -> bytes:
    """DB の assignments を正本として、拠点テンプレートから時間割 Excel を生成する。"""
    period = get_period(period_id)
    slug = resolve_location_slug(location_slug or period.location_slug)
    open_dates = open_dates_for_period(period)
    return _build_workbook_for_open_dates(open_dates, slug)


def build_schedule_workbook_for_calendar_month(
    calendar_year: int,
    month: int,
    location_slug: str | None = None,
) -> bytes:
    """指定暦月の開校日分を Excel に出力する。"""
    from services.academic_calendar import open_dates_for_calendar_month
    from services.fiscal_year_store import resolve_schedule_period_for_date

    iso_anchor = f"{calendar_year}-{month:02d}-15"
    period = resolve_schedule_period_for_date(iso_anchor)
    slug = resolve_location_slug(location_slug or period.location_slug)
    open_dates = open_dates_for_calendar_month(calendar_year, month, period.closed_dates)
    return _build_workbook_for_open_dates(open_dates, slug)


def export_calendar_month_workbook(
    calendar_year: int,
    month: int,
    location_slug: str | None = None,
) -> tuple[bytes, str]:
    """指定暦月の時間割 Excel を返す（DB 正本）。"""
    from services.fiscal_year_store import resolve_schedule_period_for_date

    iso_anchor = f"{calendar_year}-{month:02d}-15"
    period = resolve_schedule_period_for_date(iso_anchor)
    slug = resolve_location_slug(location_slug or period.location_slug)
    content = build_schedule_workbook_for_calendar_month(calendar_year, month, slug)
    filename = f"{calendar_year}年{month}月_{slug}_時間割.xlsx"
    return content, filename


def export_period_schedule_workbook(
    period_id: int,
    location_slug: str | None = None,
) -> tuple[bytes, str]:
    """DB 正本から拠点形式の時間割 Excel を返す。"""
    period = get_period(period_id)
    slug = resolve_location_slug(location_slug or period.location_slug)
    content = build_schedule_workbook_from_db(period_id, slug)
    filename = f"{period.name}_{slug}_時間割.xlsx"
    return content, filename


def save_schedule_to_output_folder(
    period_id: int,
    location_slug: str | None = None,
) -> dict:
    """DB 正本を拠点フォルダへ書き出す。output/{slug}/{period_id}_{name}/"""
    period = get_period(period_id)
    slug = resolve_location_slug(location_slug or period.location_slug)
    content = build_schedule_workbook_from_db(period_id, slug)
    out_dir = output_dir_for_period(period_id, period.name, slug)
    out_path = out_dir / f"{period.name}_時間割.xlsx"
    out_path.write_bytes(content)
    return {
        "period_id": period_id,
        "location_slug": slug,
        "output_dir": str(out_dir),
        "file_path": str(out_path),
        "message": f"{slug} 形式で {out_path} に出力しました",
    }


def export_juku_schedule_workbook(period_id: int, location_slug: str | None = None) -> bytes:
    """Google 書き込み等 — DB 正本から生成。"""
    content, _ = export_period_schedule_workbook(period_id, location_slug)
    return content
