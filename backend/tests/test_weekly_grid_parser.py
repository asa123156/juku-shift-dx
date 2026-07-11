"""週間シート（6日横並び・ブロック行数バラバラ）の日付マッピングのテスト。

実ファイル「6月度時間割【白庭台】」で、火曜以降の生徒が全員シート名の
日付（月曜）として取り込まれるバグの回帰テスト。
"""

from datetime import date, time

from services.juku_grid_parser import parse_sheet_records
from tests.weekly_grid_fixture import build_weekly_workbook

START = date(2026, 6, 29)
SLOT_BY_TIME = {time(17, 30): "4", time(19, 0): "5", time(20, 30): "6"}

ENTRIES = [
    {"day_index": 0, "time": time(17, 30), "ratio": "1：4", "seat": 1, "side": "a",
     "name": "田中花子", "grade": "小6", "subject": "算数", "teacher": "佐藤"},
    {"day_index": 1, "time": time(17, 30), "ratio": "1：4", "seat": 1, "side": "a",
     "name": "中村音娃", "grade": "小5", "subject": "国算"},
    {"day_index": 2, "time": time(19, 0), "ratio": "1：2", "seat": 3, "side": "b",
     "name": "山本次郎", "grade": "中2", "subject": "英語", "teacher": "鈴木"},
    {"day_index": 3, "time": time(20, 30), "ratio": "1：2", "seat": 5, "side": "a",
     "name": "高橋三奈", "grade": "高1", "subject": "数学", "absent": True},
    {"day_index": 5, "time": time(19, 0), "ratio": "1：4", "seat": 2, "side": "b",
     "name": "伊藤四葉", "grade": "中3", "subject": "理科", "makeup": True},
    {"day_index": 0, "time": time(20, 30), "ratio": "1：2", "seat": 11, "side": "a",
     "name": "渡辺五郎", "grade": "中1", "subject": "社会"},
    {"day_index": 4, "time": time(17, 30), "ratio": "1：2", "seat": 7, "side": "b",
     "name": "小林六実", "grade": "小4", "subject": "国語", "lesson_type": "講習"},
]


def test_weekly_sheet_maps_each_day_block_to_its_date():
    raw = build_weekly_workbook(ENTRIES, START)
    records = parse_sheet_records(raw, "6月29日", 2026)
    by_name = {r["student_name"]: r for r in records}

    assert len(records) == len(ENTRIES), "全員が抽出されること"
    for e in ENTRIES:
        r = by_name[e["name"]]
        expected_date = date.fromordinal(START.toordinal() + e["day_index"]).isoformat()
        assert r["date"] == expected_date, f"{e['name']} の日付が列ブロックと一致すること"
        assert r["slot_key"] == SLOT_BY_TIME[e["time"]]
        assert r["subject"] == e["subject"]
        assert r["grade"] == e["grade"]
        if e.get("teacher"):
            assert r["teacher_name"] == e["teacher"]


def test_daily_sheet_single_date_unaffected():
    """日付セルが1つだけ（日次シート相当）なら従来どおりシート日付を使う。"""
    entries = [ENTRIES[0]]
    raw = build_weekly_workbook(entries, START)
    # 生成器は6日分の日付を書くので、日次相当にするため2日目以降の日付を消す
    import io
    from openpyxl import load_workbook
    from tests.weekly_grid_fixture import DAY_COL_STARTS
    wb = load_workbook(io.BytesIO(raw))
    ws = wb.active
    for col in DAY_COL_STARTS[1:]:
        ws.cell(row=1, column=col + 1, value=None)
    buf = io.BytesIO()
    wb.save(buf)

    records = parse_sheet_records(buf.getvalue(), "6月29日", 2026)
    assert len(records) == 1
    assert records[0]["date"] == "2026-06-29"
