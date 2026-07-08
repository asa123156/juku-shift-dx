"""時間割表（assignments 正本）から提案書 Excel を生成する。"""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from services.entity_store import list_students
from services.period_store import get_period
from services.student_plan_store import list_plans_for_period

WEEKDAY_JA = ["月", "火", "水", "木", "金", "土", "日"]


def _format_date_label(iso_date: str) -> str:
    d = date.fromisoformat(iso_date)
    return f"{d.month}/{d.day}（{WEEKDAY_JA[d.weekday()]}）"


def _write_student_sheet(ws, proposal: dict, plans: list[dict]) -> None:
    title_font = Font(bold=True, size=14)
    header_fill = PatternFill("solid", fgColor="E8EEF4")
    header_font = Font(bold=True)

    ws["A1"] = f"{proposal['period_name']} スケジュール提案書"
    ws["A1"].font = title_font
    ws["A2"] = proposal["student_name"]
    ws["A2"].font = Font(bold=True, size=12)
    if proposal.get("grade_label"):
        ws["B2"] = proposal["grade_label"]

    plan_text = "、".join(f"{p['subject']}×{p['slot_count']}" for p in plans) if plans else "—"
    ws["A3"] = f"希望: {plan_text}"
    ws["A4"] = "※ 時間割表（教室長編集）を正本とした提案内容です"

    headers = ["日付", "コマ", "時間", "教科", "講師", "区分"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=6, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    rows = proposal["rows"]
    if not rows:
        ws.cell(row=7, column=1, value="（時間割表に割当がありません）")
    else:
        for idx, row in enumerate(rows, start=7):
            time_label = f"{row['start']}~{row['end']}" if row.get("start") else ""
            ws.cell(row=idx, column=1, value=_format_date_label(row["date"]))
            ws.cell(row=idx, column=2, value=f"{row['slot']}コマ")
            ws.cell(row=idx, column=3, value=time_label)
            ws.cell(row=idx, column=4, value=row["subject"])
            ws.cell(row=idx, column=5, value=row["teacher_name"])
            ws.cell(row=idx, column=6, value=row.get("lesson_kind") or "講習")

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 8


def export_student_proposal_xlsx(period_id: int, student_id: int) -> tuple[bytes, str]:
    proposal = build_student_proposal_rows(period_id, student_id)
    plans = list_plans_for_period(period_id).get(student_id, [])
    wb = Workbook()
    ws = wb.active
    ws.title = "提案書"
    _write_student_sheet(ws, proposal, plans)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    safe_name = proposal["student_name"].replace('"', "")
    filename = f"{proposal['period_name']}_{safe_name}_提案書.xlsx"
    return buf.getvalue(), filename


def export_period_proposals_xlsx(period_id: int) -> tuple[bytes, str]:
    """期間内の全生徒分提案書を1ブックに（シート per 生徒）。"""
    period = get_period(period_id)
    plans_by_student = list_plans_for_period(period_id)

    student_ids: set[int] = {s["id"] for s in list_students()}
    student_ids.update(plans_by_student.keys())
    from services.assignment_store import get_assignments_between

    for row in get_assignments_between(period.start_date, period.end_date):
        if (row.get("lesson_kind") or "講習") != "通常":
            student_ids.add(row["student_id"])

    wb = Workbook()
    wb.remove(wb.active)
    for sid in sorted(student_ids):
        proposal = build_student_proposal_rows(period_id, sid)
        plans = plans_by_student.get(sid, [])
        if not proposal["rows"] and not plans:
            continue
        title = proposal["student_name"][:31]
        ws = wb.create_sheet(title=title)
        _write_student_sheet(ws, proposal, plans)

    if not wb.sheetnames:
        ws = wb.create_sheet(title="提案書")
        ws["A1"] = "時間割表に割当がありません"

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    filename = f"{period.name}_提案書一括.xlsx"
    return buf.getvalue(), filename
