"""白庭台の週間シート（6日横並び・ブロック行数バラバラ）を模したテスト xlsx 生成器。

実ファイル「コピー～ 202600_6月度時間割【白庭台】」の構造を再現する:
- 行0: 日付が列 1, 21, 41, 61, 81, 101 に横並び（月〜土）
- 各時刻ブロック: 列0に時刻セル → ヘッダー行（座席/学年/氏名/科目/種別/欠席/振替/講師 ×6日分）→ 座席行
- 1:4 と 1:2 のサブブロックで座席行数が異なる（日によっても授業数はバラバラ）
- 欠席・振替はブール値セル（スプレッドシートのチェックボックス相当）
"""

from __future__ import annotations

import io
from datetime import date, time

from openpyxl import Workbook

DAY_COL_STARTS = [1, 21, 41, 61, 81, 101]  # 日付セルの列（teacher_col - 8 = 座席列と一致）
TEACHER_COLS = [9, 29, 49, 69, 89, 109]

# (時刻, [(比率ラベル, 座席数)]) — 座席数はブロックごとにバラバラにする
BLOCKS = [
    (time(17, 30), [("1：4", 2), ("1：2", 7)]),
    (time(19, 0), [("1：4", 3), ("1：2", 15)]),
    (time(20, 30), [("1：4", 2), ("1：2", 11)]),
]

HEADER = {
    -8: "座席", -7: "学年", -5: "氏名", -4: "科目", -3: "種別(通常はブランク)",
    -2: "欠席", -1: "振替\nあり", 0: "講師",
    1: "学年", 3: "氏名", 4: "科目", 5: "種別(通常はブランク)", 6: "欠席", 7: "振替\nあり",
}


def build_weekly_workbook(entries: list[dict], start: date = date(2026, 6, 29)) -> bytes:
    """entries: [{day_index, time, ratio, seat, side, name, grade, subject, lesson_type,
    teacher, absent, makeup}] を配置した週間シートを作る。"""
    wb = Workbook()
    ws = wb.active
    ws.title = f"{start.month}月{start.day}日"

    # 行0(=Excel行1): 日付横並び
    for i, col in enumerate(DAY_COL_STARTS):
        d = date.fromordinal(start.toordinal() + i)
        ws.cell(row=1, column=col + 1, value=d)

    row = 3  # Excel 1-origin
    layout = {}  # (time, ratio, seat_no) -> excel_row
    for block_time, subblocks in BLOCKS:
        for ratio, seat_count in subblocks:
            # 比率ラベル行
            for col in DAY_COL_STARTS:
                ws.cell(row=row, column=col + 1, value=ratio)
            row += 1
            # 時刻セル + ヘッダー行（実物同様、時刻はヘッダー行の列0に置くケースを再現）
            ws.cell(row=row, column=1, value=block_time)
            for tc in TEACHER_COLS:
                for off, label in HEADER.items():
                    ws.cell(row=row, column=tc + off + 1, value=label)
            row += 1
            # 座席行
            for seat_no in range(1, seat_count + 1):
                for tc in TEACHER_COLS:
                    ws.cell(row=row, column=tc - 8 + 1, value=str(seat_no))
                    # 欠席・振替はチェックボックス相当のブール値（既定 False）
                    ws.cell(row=row, column=tc - 2 + 1, value=False)
                    ws.cell(row=row, column=tc - 1 + 1, value=False)
                layout[(block_time, ratio, seat_no)] = row
                row += 1
            row += 1  # ブロック間の空行

    # 生徒を配置
    for e in entries:
        excel_row = layout[(e["time"], e["ratio"], e["seat"])]
        tc = TEACHER_COLS[e["day_index"]]
        if e["side"] == "a":
            name_col, subj_col, type_col = tc - 5, tc - 4, tc - 3
            absent_col, makeup_col = tc - 2, tc - 1
        else:
            name_col, subj_col, type_col = tc + 3, tc + 4, tc + 5
            absent_col, makeup_col = tc + 6, tc + 7
        ws.cell(row=excel_row, column=name_col + 1, value=e["name"])
        ws.cell(row=excel_row, column=name_col - 1 + 1, value=e["grade"])  # 学年ラベル(氏名の左)
        ws.cell(row=excel_row, column=subj_col + 1, value=e["subject"])
        if e.get("lesson_type"):
            ws.cell(row=excel_row, column=type_col + 1, value=e["lesson_type"])
        if e.get("teacher"):
            ws.cell(row=excel_row, column=tc + 1, value=e["teacher"])
        ws.cell(row=excel_row, column=absent_col + 1, value=bool(e.get("absent", False)))
        ws.cell(row=excel_row, column=makeup_col + 1, value=bool(e.get("makeup", False)))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
