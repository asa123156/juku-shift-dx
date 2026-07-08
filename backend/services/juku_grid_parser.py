"""白庭台形式の時間割グリッド（xlsx）を縦持ちレコードに変換する。"""

from __future__ import annotations

import io
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
from fastapi import HTTPException
from openpyxl import load_workbook

from services.entity_store import parse_grade_label
from services.location_config import load_grid_map
from services.slot_timing import generate_time_slots

SheetDate = tuple[int, int]  # month, day

SKIP_SHEET_NAME = re.compile(r"参照|講師シフト|マスタ|設定|template", re.I)
DAILY_SHEET_NAME = re.compile(r"^(\d+)月(\d+)日$")
WEEKLY_RANGE_SHEET_NAME = re.compile(r"^(\d+)月(\d+)日[～~\-](\d+)月(\d+)日$")


def load_grid_map_for_location(location_slug: str | None = None) -> dict:
    """拠点フォルダの schedule_map.json を読む（後方互換エイリアス）。"""
    return load_grid_map(location_slug)


def _cell_text(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    return "" if text.lower() in ("nan", "none") else text


def _is_header_row(row: pd.Series) -> bool:
    return sum(1 for v in row if _cell_text(v) == "氏名") >= 2


def _is_time_value(value: object) -> bool:
    return isinstance(value, time)


def _time_to_slot_key(value: time) -> str | None:
    hhmm = value.strftime("%H:%M")
    for slot in generate_time_slots():
        if slot["start"] == hhmm:
            return str(slot["slot"])
    return None


def parse_sheet_date_from_name(name: str, default_year: int) -> date | None:
    match = DAILY_SHEET_NAME.match(name.strip())
    if match is None:
        return None
    month, day = int(match.group(1)), int(match.group(2))
    return date(default_year, month, day)


def resolve_dates_for_sheet_name(name: str, default_year: int) -> list[str]:
    """シート名から対象日（複数日可）を推定する。"""
    stripped = name.strip()
    daily = parse_sheet_date_from_name(stripped, default_year)
    if daily is not None:
        return [daily.isoformat()]

    weekly = WEEKLY_RANGE_SHEET_NAME.match(stripped)
    if weekly is not None:
        start = date(default_year, int(weekly.group(1)), int(weekly.group(2)))
        end = date(default_year, int(weekly.group(3)), int(weekly.group(4)))
        dates: list[str] = []
        cur = start
        while cur <= end:
            dates.append(cur.isoformat())
            cur += timedelta(days=1)
        return dates

    return []


def _scan_dates_in_rows(rows: list[tuple], default_year: int) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()
    for row in rows[:300]:
        if not row:
            continue
        for val in row:
            parsed = _parse_sheet_date_from_cell(val, default_year)
            if parsed is None:
                text = _cell_text(val)
                md = DAILY_SHEET_NAME.match(text)
                if md:
                    parsed = date(default_year, int(md.group(1)), int(md.group(2)))
            if parsed is None:
                continue
            iso = parsed.isoformat()
            if iso not in seen:
                seen.add(iso)
                found.append(iso)
    return found


def resolve_dates_for_sheet(
    sheet_name: str,
    rows: list[tuple],
    default_year: int,
    cfg: dict | None = None,
) -> list[str]:
    """時間割シートがカバーする日付一覧（推定）。"""
    from_name = resolve_dates_for_sheet_name(sheet_name, default_year)
    if from_name:
        return from_name

    grid_cfg = cfg or load_grid_map()
    dr, dc = grid_cfg["date_header_row"], grid_cfg["date_header_col"]
    if dr < len(rows) and dc < len(rows[dr]):
        cell_date = _parse_sheet_date_from_cell(rows[dr][dc], default_year)
        if cell_date is not None:
            return [cell_date.isoformat()]

    scanned = _scan_dates_in_rows(rows, default_year)
    if scanned:
        return scanned

    return []


def is_schedule_grid_rows(rows: list[tuple]) -> bool:
    """時間列 + 氏名ヘッダーがあるシートか（参照・講師シフト等は除外）。"""
    has_time = False
    name_headers = 0
    for row in rows[:250]:
        if not row:
            continue
        if isinstance(row[0], time):
            has_time = True
        name_headers = max(name_headers, sum(1 for v in row if _cell_text(v) == "氏名"))
        if has_time and name_headers >= 2:
            return True
    return False


def list_schedule_sheet_names(raw: bytes) -> list[str]:
    """時間割グリッドを含むシート名のみ返す（参照・講師シフト等は除外）。"""
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Excel ファイルを読み込めません") from exc

    names: list[str] = []
    for name in wb.sheetnames:
        if SKIP_SHEET_NAME.search(name):
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        if is_schedule_grid_rows(rows):
            names.append(name)
    wb.close()
    return names


def _parse_sheet_date_from_name(name: str, default_year: int) -> date | None:
    return parse_sheet_date_from_name(name, default_year)


def _parse_sheet_date_from_cell(value: object, default_year: int) -> date | None:
    if isinstance(value, (datetime, date)):
        dt = value.date() if isinstance(value, datetime) else value
        return dt.replace(year=default_year) if default_year else dt
    text = _cell_text(value)
    if not text:
        return None
    parsed = pd.to_datetime(text, errors="coerce")
    if pd.isna(parsed):
        return None
    dt = parsed.date()
    return dt.replace(year=default_year)


def _teacher_col_for_name_col(name_col: int, teacher_columns: list[int]) -> int | None:
    for tc in teacher_columns:
        if tc - 8 <= name_col <= tc + 5:
            return tc
    return None


def _name_columns_in_header(header: pd.Series) -> list[int]:
    return [i for i, v in enumerate(header) if _cell_text(v) == "氏名"]


def _read_grade_label(row: pd.Series, name_col: int, grade_label_shift: int) -> str:
    for offset in (grade_label_shift, -2, -1, 1):
        col = name_col + offset
        if col < 0 or col >= len(row):
            continue
        text = _cell_text(row.iloc[col])
        if not text:
            continue
        try:
            parse_grade_label(text)
            return text
        except ValueError:
            continue
    return ""


def _looks_like_person_name(text: str) -> bool:
    if not text or len(text) < 2:
        return False
    if text.isdigit():
        return False
    blocked = {"氏名", "講師", "科目", "学年", "座席", "月曜日", "予定", "実際", "欠席"}
    if text in blocked:
        return False
    return any("\u4e00" <= ch <= "\u9fff" for ch in text)


def parse_grid_dataframe(
    df: pd.DataFrame,
    iso_date: str,
    grid_map: dict | None = None,
    *,
    detect_dates_in_rows: bool = False,
    period_year: int | None = None,
) -> list[dict[str, Any]]:
    """1日分（または週次）シートの DataFrame から縦持ちレコードを抽出する。"""
    cfg = grid_map or load_grid_map()
    teacher_columns: list[int] = cfg["teacher_columns"]
    offsets = cfg["offsets_from_teacher"]
    grade_shift = int(cfg.get("grade_label_col_shift", -1))

    records: list[dict[str, Any]] = []
    current_slot: str | None = None
    header_row_idx: int | None = None
    current_date = iso_date

    for row_idx in range(len(df)):
        row = df.iloc[row_idx]
        if detect_dates_in_rows and period_year is not None:
            for val in row:
                parsed = _parse_sheet_date_from_cell(val, period_year)
                if parsed is not None:
                    current_date = parsed.isoformat()
                    break
                text = _cell_text(val)
                md = DAILY_SHEET_NAME.match(text)
                if md:
                    current_date = date(period_year, int(md.group(1)), int(md.group(2))).isoformat()
                    break

        time_val = row.iloc[0] if len(row) > 0 else None
        if _is_time_value(time_val):
            slot_key = _time_to_slot_key(time_val)
            if slot_key is not None:
                current_slot = slot_key
            if _is_header_row(row):
                header_row_idx = row_idx
            continue

        if _is_header_row(row):
            header_row_idx = row_idx
            continue

        if header_row_idx is None or current_slot is None:
            continue

        header = df.iloc[header_row_idx]
        for name_col in _name_columns_in_header(header):
            teacher_col = _teacher_col_for_name_col(name_col, teacher_columns)
            if teacher_col is None:
                continue

            is_b = name_col > teacher_col
            prefix = "b" if is_b else "a"
            student_name = _cell_text(row.iloc[name_col])
            if not _looks_like_person_name(student_name):
                continue

            grade = _read_grade_label(row, name_col, grade_shift)
            subject_col = name_col + 1
            type_col = name_col + 2
            subject = _cell_text(row.iloc[subject_col]) if subject_col < len(row) else ""
            lesson_type = _cell_text(row.iloc[type_col]) if type_col < len(row) else ""
            teacher_name = _cell_text(row.iloc[teacher_col]) if teacher_col < len(row) else ""

            seat_col = teacher_col + offsets["seat"]
            seat = _cell_text(row.iloc[seat_col]) if 0 <= seat_col < len(row) else ""

            records.append(
                {
                    "date": current_date,
                    "slot_key": current_slot,
                    "seat": seat,
                    "student_name": student_name,
                    "grade": grade,
                    "subject": subject,
                    "lesson_type": lesson_type,
                    "teacher_name": teacher_name,
                    "teacher_col": teacher_col,
                    "name_col": name_col,
                    "row_idx": row_idx,
                    "student_side": prefix,
                }
            )

    return records


def parse_sheet_records(
    raw: bytes,
    sheet_name: str,
    period_year: int,
    location_slug: str | None = None,
) -> list[dict[str, Any]]:
    """指定シートから縦持ちレコードを抽出する（週次シートは行内日付も参照）。"""
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Excel ファイルを読み込めません") from exc

    if sheet_name not in wb.sheetnames:
        wb.close()
        raise HTTPException(status_code=400, detail=f"シート「{sheet_name}」が見つかりません")

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []

    cfg = load_grid_map(location_slug)
    dates = resolve_dates_for_sheet(sheet_name, rows, period_year, cfg)
    fallback = dates[0] if dates else date(period_year, 1, 1).isoformat()
    multi_day = len(dates) > 1 or not dates

    df = pd.DataFrame(rows)
    return parse_grid_dataframe(
        df,
        fallback,
        cfg,
        detect_dates_in_rows=multi_day,
        period_year=period_year,
    )


def read_workbook_sheet(
    raw: bytes,
    iso_date: str,
    period_year: int,
    sheet_name: str | None = None,
    location_slug: str | None = None,
) -> tuple[list[dict[str, Any]], str]:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Excel ファイルを読み込めません") from exc

    target_sheet = sheet_name
    if target_sheet is None:
        target = date.fromisoformat(iso_date).replace(year=period_year)
        expected = f"{target.month}月{target.day}日"
        if expected in wb.sheetnames:
            target_sheet = expected
        else:
            daily = [s for s in wb.sheetnames if DAILY_SHEET_NAME.match(s)]
            if not daily:
                raise HTTPException(status_code=400, detail="日次シートが見つかりません")
            target_sheet = daily[0]

    if target_sheet not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"シート「{target_sheet}」が見つかりません")

    ws = wb[target_sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise HTTPException(status_code=400, detail="シートが空です")

    df = pd.DataFrame(rows)
    cfg = load_grid_map(location_slug)
    resolved_date = iso_date
    if sheet_name is None:
        parsed = _parse_sheet_date_from_name(target_sheet, period_year)
        if parsed is not None:
            resolved_date = parsed.isoformat()
        else:
            dr, dc = cfg["date_header_row"], cfg["date_header_col"]
            if dr < len(rows) and dc < len(rows[dr]):
                cell_date = _parse_sheet_date_from_cell(rows[dr][dc], period_year)
                if cell_date is not None:
                    resolved_date = cell_date.isoformat()

    records = parse_grid_dataframe(df, resolved_date, cfg)
    return records, target_sheet


def parse_juku_grid_xlsx(
    raw: bytes,
    iso_date: str,
    period_year: int,
    sheet_name: str | None = None,
    location_slug: str | None = None,
) -> list[dict[str, Any]]:
    records, _ = read_workbook_sheet(raw, iso_date, period_year, sheet_name, location_slug)
    return records
