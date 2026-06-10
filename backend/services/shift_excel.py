import csv
import io
from typing import Iterator

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from services.period_store import get_period, iter_dates, set_period_base_slot

_REQUIRED = {"role", "entity_id", "date", "slot", "symbol"}


def _normalize_header(name: str) -> str:
    return name.strip().lower().replace(" ", "_")


def _parse_rows(rows: Iterator[dict[str, str]]) -> list[dict[str, str]]:
    parsed: list[dict[str, str]] = []
    for line_no, raw in enumerate(rows, start=2):
        row = {_normalize_header(k): (v or "").strip() for k, v in raw.items() if k}
        role = row.get("role", "").lower()
        if role not in ("teacher", "student"):
            raise HTTPException(status_code=400, detail=f"Line {line_no}: role must be teacher or student")

        iso_date = row.get("date", "")
        try:
            entity_id = int(row.get("entity_id", ""))
            slot = str(int(row.get("slot", "")))
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=f"Line {line_no}: invalid entity_id or slot") from exc

        if slot not in ("1", "2", "3", "4"):
            raise HTTPException(status_code=400, detail=f"Line {line_no}: slot must be 1-4")

        symbol = row.get("symbol", "")
        parsed.append(
            {
                "line_no": str(line_no),
                "role": role,
                "date": iso_date,
                "entity_id": str(entity_id),
                "slot": slot,
                "symbol": symbol,
            }
        )
    return parsed


def _import_parsed_rows(period_id: int, rows: list[dict[str, str]]) -> int:
    period = get_period(period_id)
    if period.status not in ("DRAFT", "COLLECTING"):
        raise HTTPException(status_code=409, detail="Import only allowed in DRAFT or COLLECTING")

    allowed_dates = set(iter_dates(period.start_date, period.end_date))
    count = 0
    for row in rows:
        line_no = row.get("line_no", "?")
        iso_date = row["date"]
        if iso_date not in allowed_dates:
            raise HTTPException(status_code=400, detail=f"Line {line_no}: date outside period")

        symbol = row["symbol"]
        if symbol != "◎":
            raise HTTPException(status_code=400, detail=f"Line {line_no}: import symbol must be ◎")

        set_period_base_slot(
            period_id,
            row["role"],
            int(row["entity_id"]),
            iso_date,
            row["slot"],
            symbol,
        )
        count += 1
    return count


def import_shift_excel_csv(period_id: int, content: str) -> int:
    reader = csv.DictReader(io.StringIO(content))
    if reader.fieldnames is None:
        raise HTTPException(status_code=400, detail="CSV header missing")

    headers = {_normalize_header(h) for h in reader.fieldnames if h}
    if not _REQUIRED.issubset(headers):
        missing = ", ".join(sorted(_REQUIRED - headers))
        raise HTTPException(status_code=400, detail=f"Missing columns: {missing}")

    rows = _parse_rows(reader)
    return _import_parsed_rows(period_id, rows)


def import_shift_excel_xlsx(period_id: int, raw: bytes) -> int:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Excel ファイルを読み込めません") from exc

    ws = wb.active
    if ws is None:
        raise HTTPException(status_code=400, detail="Excel シートが空です")

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise HTTPException(status_code=400, detail="Excel ヘッダー行がありません") from exc

    headers = [_normalize_header(str(h or "")) for h in header_row]
    if not _REQUIRED.issubset(set(headers)):
        missing = ", ".join(sorted(_REQUIRED - set(headers)))
        raise HTTPException(status_code=400, detail=f"Missing columns: {missing}")

    dict_rows: list[dict[str, str]] = []
    for values in rows_iter:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        row_dict = {}
        for idx, key in enumerate(headers):
            if key:
                row_dict[key] = str(values[idx]).strip() if idx < len(values) and values[idx] is not None else ""
        dict_rows.append(row_dict)

    parsed = _parse_rows(dict_rows)
    return _import_parsed_rows(period_id, parsed)


def export_shift_excel_xlsx(period_id: int) -> bytes:
    from services.schedule_service import build_merged_slots_for_export

    period = get_period(period_id)
    dates = iter_dates(period.start_date, period.end_date)
    wb = Workbook()
    ws = wb.active
    ws.title = "shifts"
    ws.append(["role", "entity_id", "date", "slot", "symbol"])

    for iso_date in dates:
        for teacher_id in (1, 2, 3):
            slots = build_merged_slots_for_export("teacher", teacher_id, period.id, iso_date)
            for slot in ("1", "2", "3", "4"):
                ws.append(["teacher", teacher_id, iso_date, slot, slots.get(slot, "")])
        for student_id in (1, 2, 3):
            slots = build_merged_slots_for_export("student", student_id, period.id, iso_date)
            for slot in ("1", "2", "3", "4"):
                ws.append(["student", student_id, iso_date, slot, slots.get(slot, "")])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
